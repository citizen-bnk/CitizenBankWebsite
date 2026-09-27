"""Document text extraction service for various file types."""

from typing import Optional
import io
import PyPDF2
import docx


class DocumentExtractor:
    """Extract text content from various document types."""

    def extract_text(self, file_content: bytes, file_type: str, file_name: str) -> str:
        """Extract text from document based on file type."""
        try:
            if 'pdf' in file_type.lower() or file_name.lower().endswith('.pdf'):
                return self._extract_from_pdf(file_content)
            elif 'word' in file_type.lower() or file_name.lower().endswith(('.doc', '.docx')):
                return self._extract_from_docx(file_content)
            elif 'text' in file_type.lower() or file_name.lower().endswith('.txt'):
                return file_content.decode('utf-8', errors='ignore')
            else:
                # For unknown types, try to decode as text
                return file_content.decode('utf-8', errors='ignore')
        except Exception as e:
            print(f"Error extracting text from {file_name}: {str(e)}")
            return ""

    def _extract_from_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF file."""
        try:
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text_parts = []
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            
            return "\n".join(text_parts)
        except Exception as e:
            print(f"PDF extraction error: {str(e)}")
            return ""

    def _extract_from_docx(self, file_content: bytes) -> str:
        """Extract text from Word document."""
        try:
            doc_file = io.BytesIO(file_content)
            doc = docx.Document(doc_file)
            
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text)
            
            return "\n".join(text_parts)
        except Exception as e:
            print(f"DOCX extraction error: {str(e)}")
            return ""

    def get_text_preview(self, text: str, max_length: int = 500) -> str:
        """Get preview of extracted text."""
        if len(text) <= max_length:
            return text
        return text[:max_length] + "..."
